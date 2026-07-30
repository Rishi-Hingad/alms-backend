import frappe
from openpyxl import load_workbook

@frappe.whitelist()
def process_vendor_excel(file_url):
    try:
        file_doc = frappe.get_doc("File", {"file_url": file_url})
        file_path = file_doc.get_full_path()

        wb = load_workbook(file_path, data_only=True)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]
        row_cells = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), None)
        row = dict(zip(headers, row_cells)) if row_cells else {}

        return {
            "financed_amount": row.get("Financed Amount"),
            "location": row.get("Location"),
            "variant": row.get("Variant"),
        }

    except Exception:
        frappe.log_error(frappe.get_traceback(), "Excel Processing Error")
        frappe.throw("Failed to process Excel file")